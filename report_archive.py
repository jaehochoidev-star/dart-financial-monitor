"""Daily report storage and a JavaScript-free GitHub Pages archive."""

import json
import shutil
from datetime import datetime, timedelta, timezone
from html import escape, unescape
from pathlib import Path
from openpyxl import load_workbook

KST = timezone(timedelta(hours=9))


def plain_message(message):
    # The shared message only uses these Telegram formatting tags.
    for tag in ("<b>", "</b>", "<code>", "</code>"):
        message = message.replace(tag, "")
    return unescape(message)


def archive_report(message, base_date, total_count, matched_count,
                   attachment_path=None, root="reports", run_date=None):
    now = datetime.now(KST)
    day = run_date or now.strftime("%Y-%m-%d")
    datetime.strptime(day, "%Y-%m-%d")
    directory = Path(root) / day
    directory.mkdir(parents=True, exist_ok=True)
    attachment = None
    if attachment_path:
        attachment = "report.xlsx"
        shutil.copyfile(attachment_path, directory / attachment)
    record = {
        "date": day, "updated_at": now.isoformat(), "base_date": base_date,
        "total_count": total_count, "matched_count": matched_count,
        "message": plain_message(message), "attachment": attachment,
    }
    temporary = directory / "report.json.tmp"
    temporary.write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(directory / "report.json")


def render_financial_table(filepath, day):
    """Render the existing report workbook without changing its numeric values."""
    workbook = load_workbook(filepath, read_only=True, data_only=True)
    try:
        sheet = workbook["증가 기업 목록"]
        rows = sheet.iter_rows(min_row=3, max_col=15, values_only=True)
        headers = next(rows)
        headings = ''.join(f'<th scope="col">{escape(str(value or ""))}</th>' for value in headers)
        body = []
        for row in rows:
            if all(value is None for value in row):
                continue
            cells = []
            for index, value in enumerate(row):
                if value is None:
                    label = "—"
                elif index == 0:
                    label = str(value).zfill(6)
                elif index in (8, 11, 14) and isinstance(value, (int, float)):
                    # The workbook stores 42.88 for 42.88%, not 0.4288.
                    label = f"{value:,.2f}%"
                elif index >= 6 and isinstance(value, (int, float)):
                    label = f"{value:,.0f}"
                else:
                    label = str(value)
                classes = ['number'] if index >= 6 else []
                if index in (8, 11, 14) and isinstance(value, (int, float)):
                    classes.append('positive' if value > 0 else 'negative' if value < 0 else 'neutral')
                alignment = f' class="{" ".join(classes)}"' if classes else ''
                cells.append(f'<td{alignment}>{escape(label)}</td>')
            body.append('<tr>' + ''.join(cells) + '</tr>')
        if not body:
            return '<p>상세 내역에 표시할 기업이 없습니다.</p>'
        return (f'<h3 id="table-{day}">기업별 상세 내역</h3>'
                '<p class="meta">금액 단위: 원 · 증감률: % · 표를 좌우로 스크롤하면 모든 항목을 볼 수 있습니다.</p>'
                f'<div class="table-scroll" role="region" aria-labelledby="table-{day}" tabindex="0">'
                f'<table><caption>{day} 기업별 재무 비교</caption><thead>'
                '<tr class="groups"><th colspan="6" scope="colgroup">기업 정보</th>'
                '<th colspan="3" scope="colgroup">매출액</th><th colspan="3" scope="colgroup">영업이익</th>'
                '<th colspan="3" scope="colgroup">당기순이익</th></tr>'
                f'<tr>{headings}</tr></thead>'
                '<tbody>' + ''.join(body) + '</tbody></table></div>')
    finally:
        workbook.close()


def build_site(root="reports", output="site"):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    cards_by_month = {}
    for source in sorted(Path(root).glob("????-??-??/report.json"), reverse=True):
        record = json.loads(source.read_text(encoding="utf-8"))
        day = source.parent.name
        datetime.strptime(day, "%Y-%m-%d")
        month = day[:7]
        link = ""
        detail_table = ""
        if record["attachment"]:
            target = output / "downloads" / day
            target.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source.parent / "report.xlsx", target / "report.xlsx")
            link = f'<a class="download" href="downloads/{day}/report.xlsx" download>엑셀로 저장 ↓</a>'
            detail_table = render_financial_table(source.parent / "report.xlsx", day)
        base = record["base_date"]
        base = f"{base[:4]}-{base[4:6]}-{base[6:]}" if len(base) == 8 else "해당 공시 없음"
        summary = record['message'].replace(
            "📎 상세 내역은 첨부 엑셀 파일을 확인하세요.",
            "📋 상세 내역은 아래 표에서 확인하세요.")
        cards_by_month.setdefault(month, []).append(f'''<article>
<div class="heading"><h2><time datetime="{day}">{day}</time></h2>
<span class="badge">조건 충족 {int(record['matched_count'])}개</span></div>
<p class="meta">공시 기준일 {escape(base)} · 조회 기업 {int(record['total_count'])}개</p>
<details class="report-summary"><summary>알림 요약 보기</summary><pre>{escape(summary)}</pre></details>
{detail_table}{link}</article>''')
    template = '''<!doctype html>
<html lang="ko"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="description" content="DART 재무성장 기업 모니터링의 날짜별 결과">
<title>DART | __MONTH_TITLE__</title>
<style>
:root{color-scheme:light;font-family:system-ui,-apple-system,"Malgun Gothic",sans-serif;color:#172b3a;background:#f3f5f7}
*{box-sizing:border-box}body{margin:0;border-top:5px solid #166b60}main{max-width:1400px;margin:auto;padding:48px 24px}
header{margin-bottom:32px}.eyebrow{color:#166b60;font-size:13px;font-weight:750;letter-spacing:.12em}
h1{font-size:32px;letter-spacing:-.06em;margin:12px 0}header p,.meta,footer{color:#63717d;line-height:1.7}
article{background:white;border:1px solid #dce3e7;border-radius:14px;padding:28px;margin:0 0 20px}
.heading{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap}h2{font-size:23px;margin:0}
.badge{background:#e4f3ed;color:#176348;padding:7px 12px;border-radius:30px;font-size:13px;font-weight:650}
.meta{font-size:13px;margin:10px 0 20px}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;line-height:1.85;font-size:14px}
a{display:inline-block;color:#126458;font-weight:650;text-underline-offset:4px;margin-top:12px}a:focus-visible{outline:3px solid #166b60;outline-offset:5px}
.months{display:flex;gap:8px;overflow-x:auto;padding:6px 3px 14px;margin-bottom:16px}
.months a{flex-shrink:0;margin:0;padding:10px 16px;border:1px solid #cbd8d5;border-radius:8px;text-decoration:none;background:white}
.months a[aria-current="page"]{background:#166b60;color:white;border-color:#166b60}
.month-title{font-size:19px;margin:0 0 20px}.month-title span{font-size:13px;font-weight:400;color:#63717d;margin-left:10px}
h3{font-size:17px;margin:24px 0 8px}.table-scroll{max-width:100%;overflow-x:auto;border:1px solid #dce3e7;border-radius:8px}
.table-scroll:focus-visible{outline:3px solid #166b60;outline-offset:3px}table{border-collapse:collapse;width:100%;font-size:13px;font-variant-numeric:tabular-nums}
caption{text-align:left;padding:12px;font-weight:650}th,td{padding:12px;white-space:nowrap;text-align:left;border-bottom:1px solid #e3e9ec}
th{background:#eaf2ef;color:#285648}td.number{text-align:right}tbody tr:nth-child(even){background:#f7f9fa}tbody tr:last-child td{border-bottom:0}
/* Calm report surface, with a clear hierarchy and readable financial columns. */
:root{color:#202f42;background:#f4f6fa;font-size:15px;line-height:1.6}
body{border-top:4px solid #164e63;background:linear-gradient(180deg,#eaf0f5 0,#f4f6fa 340px)}
main{max-width:1500px;padding:40px 36px 56px}
header{padding:12px 0 28px;margin-bottom:20px;border-bottom:1px solid #d3dce5}
.eyebrow{font-size:11px;letter-spacing:.2em;color:#42617c;font-weight:750}
h1{font-size:clamp(27px,3vw,38px);line-height:1.3;letter-spacing:-.045em;margin:12px 0 14px;color:#152c43}
header p{font-size:14px;max-width:660px;margin:0;color:#526477;line-height:1.85}
.months{gap:8px;margin:0 0 24px;padding:4px 3px 12px;scrollbar-width:thin}
.months a{padding:10px 20px;min-height:44px;border-radius:10px;border-color:#d5dfe8;color:#455b70;font-size:14px}
.months a:hover{background:#edf3f7;border-color:#95adbf}
.months a[aria-current="page"]{background:#193e57;border-color:#193e57;box-shadow:0 3px 8px #17384d18;color:white}
.month-title{font-size:22px;letter-spacing:-.025em;margin:0 0 20px}
.month-title span{display:inline-block;font-size:13px;color:#607184}
article{border:1px solid #dbe3eb;border-radius:16px;padding:28px 30px;margin-bottom:28px;box-shadow:0 4px 20px #253c5605}
.heading{gap:16px}.heading h2{font-size:25px;font-weight:750;letter-spacing:-.035em;font-variant-numeric:tabular-nums}
.badge{font-size:13px;border:1px solid #c4e5d7;background:#edf8f2;color:#206044;padding:7px 13px}
.meta{color:#5f7082;line-height:1.8;font-size:13px;margin:8px 0 18px}
.report-summary{margin:20px 0;border-top:1px solid #e4eaf0;border-bottom:1px solid #e4eaf0}
.report-summary summary{cursor:pointer;padding:13px 0;font-size:13px;font-weight:650;color:#466078;min-height:44px}
.report-summary pre{background:#f5f8fb;border-radius:8px;padding:18px 20px;margin:0 0 16px;font-size:14px;line-height:1.95;color:#3e5368}
h3{font-size:17px;letter-spacing:-.02em;margin-top:24px;color:#263e54}
.table-scroll{border-color:#d8e2eb;border-radius:10px;scrollbar-width:thin;scrollbar-color:#9eb1c1 #edf2f6}
table{border-collapse:separate;border-spacing:0;font-size:14px;line-height:1.6}
caption{font-size:12px;color:#586d80;padding:12px 16px;background:#f8fafc}
th,td{padding:14px 16px;border-bottom:1px solid #e5ebf1;min-width:105px;background:white}
thead th{font-size:12px;font-weight:650;color:#486177;background:#edf3f7}
.groups th{font-size:13px;text-align:center;background:#233f55;color:#f6f9fc;border-color:#3d566b;padding:10px 16px}
.groups th+th{border-left:1px solid #546b7e}
td:nth-child(7),td:nth-child(10),td:nth-child(13),thead tr:last-child th:nth-child(7),thead tr:last-child th:nth-child(10),thead tr:last-child th:nth-child(13){border-left:2px solid #dbe5ed}
tbody tr:nth-child(even) td{background:#f7f9fc}tbody tr:hover td{background:#edf4f8}
tbody td:nth-child(2){font-weight:700;color:#213e55;min-width:200px}
tbody td:first-child{font-size:12px;color:#64778a;font-family:ui-monospace,Consolas,monospace}
thead tr:last-child th:first-child,tbody td:first-child{position:sticky;left:0;width:105px;min-width:105px;max-width:105px;z-index:1}
thead tr:last-child th:nth-child(2),tbody td:nth-child(2){position:sticky;left:105px;z-index:1;border-right:1px solid #cbd8e4;box-shadow:5px 0 8px #233f5508}
td.positive{color:#16664b;font-weight:750}td.negative{color:#b03542;font-weight:750}td.neutral{color:#53677b}
.download{font-size:13px;text-decoration:none;border:1px solid #cedbe5;padding:8px 14px;min-height:40px;border-radius:8px;margin-top:18px;color:#3d596e}
.download:hover{background:#eef4f8}a:focus-visible,summary:focus-visible,.table-scroll:focus-visible{outline:3px solid #237ca3;outline-offset:3px}
footer{font-size:12px;margin-top:36px;padding-top:20px;border-top:1px solid #dce4ec;color:#64768a;line-height:1.9}
@media(max-width:640px){main{padding:24px 14px 36px}header{padding-top:4px;padding-bottom:22px}.heading h2{font-size:22px}article{padding:20px 16px;border-radius:12px;margin-bottom:20px}.badge{font-size:12px;padding:6px 10px}.month-title{font-size:20px}.months a{padding:9px 15px}th,td{padding:12px;font-size:13px}.report-summary pre{padding:14px;font-size:13px}thead tr:last-child th:first-child,tbody td:first-child{position:static}thead tr:last-child th:nth-child(2),tbody td:nth-child(2){left:0;min-width:145px;max-width:160px;white-space:normal}.meta{font-size:12px}}
@media print{body{background:white}main{padding:0;max-width:none}.months,.download,.report-summary{display:none}article{box-shadow:none}.table-scroll{overflow:visible}table{font-size:8px}th,td{padding:4px;min-width:0!important;position:static!important}}
</style></head><body><main><header><div class="eyebrow">DART / DAILY REPORT</div>
<h1>재무성장 기업 모니터링</h1><p>월을 선택하면 해당 월의 결과를 최신 날짜부터 볼 수 있습니다.<br>매출액 · 영업이익 · 당기순이익이 모두 증가한 기업을 기록합니다.</p>
</header>__MONTH_CONTENT__<footer>날짜는 한국시간 실행일 기준입니다. 휴일에는 이전 공시일의 결과가 반복될 수 있습니다.<br>같은 날 다시 실행하면 해당 날짜의 결과가 갱신됩니다.</footer></main></body></html>'''
    months = sorted(cards_by_month, reverse=True)
    for month in months or [None]:
        title = f"{month[:4]}년 {int(month[5:])}월" if month else "일별 모니터링 기록"
        tabs = []
        for target_month in months:
            current = ' aria-current="page"' if target_month == month else ''
            label = f"{target_month[:4]}년 {int(target_month[5:])}월"
            tabs.append(f'<a href="{target_month}.html"{current}>{label}</a>')
        if month:
            content = ('<nav class="months" aria-label="월별 결과">' + ''.join(tabs) + '</nav>'
                       + f'<h2 class="month-title">{title}<span>{len(cards_by_month[month])}일의 기록</span></h2>'
                       + '\n'.join(cards_by_month[month]))
        else:
            content = '<article><h2>아직 저장된 결과가 없습니다.</h2><p>첫 모니터링 실행 후 날짜별 결과가 여기에 쌓입니다.</p></article>'
        page = template.replace('__MONTH_TITLE__', title).replace('__MONTH_CONTENT__', content)
        if month:
            (output / f"{month}.html").write_text(page, encoding="utf-8")
        if not month or month == months[0]:
            (output / "index.html").write_text(page, encoding="utf-8")
    (output / ".nojekyll").touch()


if __name__ == "__main__":
    build_site()
